#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""make_seed_pack.py - 初期ナレッジ(シードパック)を作る。

なぜ専用ツールなのか
--------------------
配布した .xlsm を開いた瞬間、利用者は「で、何をすればいいの?」で止まる。
本棚は空、質問しても資料が無い。そこで終わる。

そこで実物の社内資料をあらかじめ積んで出す。ただし「それっぽい一般知識」は
絶対に入れない。損保で初手の回答が微妙にズレたら、その時点で信用は終わる。
入れるのは実物だけ、出典は実物のページ、答えられないことは答えない。

このツールは PDF から pack形式(.xlsx)を作る。実行時の取込経路(modExtractor→
modChunker)とは別に、ここで「普段より丁寧に」切る:

  1. 文字化けページを落とす
     約款PDFの表紙・裏の連絡先ページは装飾フォント(ToUnicode無し)で
     作られていることがあり、抽出すると制御文字とギリシャ文字の羅列になる。
     本文は正常なので総文字数チェックには引っかからない。ページ単位で捨てる。

  2. 「第N条」を境界にして切る
     約款は条文が意味の単位。固定長で切ると1つの条が2つに割れ、
     「第12条は?」に対して条文の後半だけが当たる、という最悪の外し方をする。
     条の途中で切らない。長い条だけは項(1. 2. 3.)で分割する。
     ガイドラインは番号見出し(1. 1-1. (1))を境界にする。

  3. 1チャンクに必ず出所の見出しを残す
     「第12条(保険金を支払わない場合)」という見出しがチャンク本文の先頭に
     入っていれば、AIは何条の話かを取り違えない。切ったあとに見出しを
     復元して付け直す。

  4. 目次・空白ページを捨てる
     ドットリーダーだらけの目次は検索の邪魔にしかならない。

使い方
------
  python3 tools/make_seed_pack.py <出力先.xlsx> <PDF>[:表示名] ...
  python3 tools/make_seed_pack.py --self-test
      Fnv1a64Hex/NormalizeForHash が modUtil(VBA)側と同じ値を返すことを
      ゴールデン値で確認する(R12-8。詳細は self_test() 直前のコメント)。

出力は modPack と同じ3シート(pack_meta/pack_chunks/pack_vectors)。
pack_vectors は空で出す。ベクトルは社内AIリボンでしか作れないため、
配布前に一度、リボンのあるPCでこのパックを取り込み→再書き出しして
ベクトル入りにする(手順は docs/60_初期ナレッジの作り方.md)。
"""
from __future__ import annotations

import os
import re
import sys
import unicodedata
from dataclasses import dataclass

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

# --self-test(R12-8: Fnv1a64Hex/NormalizeForHashのゴールデン値照合)はPDF/xlsxに
# 一切触れない純粋な文字列アルゴリズム検証なので、pypdf/openpyxlが未インストール
# の環境(CIやゴールデン値の単体確認だけしたい場合)でも実行できるようにする。
_SELF_TEST_ONLY = "--self-test" in sys.argv[1:]

if not _SELF_TEST_ONLY:
    try:
        import pypdf
    except ImportError:
        sys.exit("pypdf が必要です: pip install pypdf")

    try:
        from openpyxl import Workbook
    except ImportError:
        sys.exit("openpyxl が必要です: pip install openpyxl")

    import pdf_layout

PACK_FORMAT_VERSION = 1
EMBED_DIM = 1536

# 1チャンクの狙い。実行時の既定(700)より大きめに取る。条文は分割するより
# 丸ごと1チャンクに収めたほうが、回答も出典も正確になる。
TARGET_CHARS = 1100
MAX_CHARS = 1800          # modChunker と同じ上限に合わせる
MIN_CHARS = 60            # これ未満の断片は捨てる(見出しだけの行など)

# 座標ベースの段組復元を使うか(--layout で有効)。既定オフの理由は
# extract_pages の説明を参照。
USE_LAYOUT = False

# 条見出し。必ず行頭が「第N条」であること。「(1)第20条…」のような
# 本文中の相互参照を見出しと誤認しないための厳格版。
ARTICLE_RE = re.compile(r"^(第\d+条(?:の\d+)?)\s*[(（]?([^)）\n]{0,40})?[)）]?")
# 節・番号見出し。文章(句点で終わる/敬体)は見出しではない。
HEADING_RE = re.compile(r"^((?:\d+[-－.．]){0,3}\d+[.．)）]|[◆■●▲]|第[一二三四五六七八九十]+[章節])\s*\S")


def is_heading_line(line: str) -> bool:
    """見出しらしさの判定。長い行・文章・条文の相互参照は見出しにしない。"""
    t = line.strip()
    if not t or len(t) > 40:
        return False
    if t.endswith(("。", "、", "，")):
        return False
    if "条(" in t or "条（" in t:      # 「…第20条(保険料の返還)…」は本文中の参照
        return False
    if any(w in t for w in ("します", "ます。", "とします", "できます", "ください")):
        return False
    return bool(HEADING_RE.match(t))


# ---------------------------------------------------------------- 抽出


def garble_ratio(text: str) -> float:
    """日本語の業務文書に出ない文字(制御文字/ギリシャ/キリル)の比率。"""
    bad = tot = 0
    for ch in text:
        if ch.isspace():
            continue
        tot += 1
        o = ord(ch)
        if o < 32:
            bad += 1
        elif 0x370 <= o <= 0x3FF or 0x400 <= o <= 0x52F:
            bad += 1
    return (bad / tot) if tot else 0.0


def looks_like_toc(text: str) -> bool:
    """目次ページか。ドットリーダーが本文の3割を超えたら目次とみなす。"""
    stripped = [c for c in text if not c.isspace()]
    if not stripped:
        return False
    dots = sum(1 for c in stripped if c in ".．…・")
    return dots / len(stripped) > 0.30


# PDFの字詰めのせいで「第2 0条」「1 2 0 日」のように数字が空白で割れる。
# そのままだと利用者が「第20条」と聞いても検索に一致しない。実物の約款で
# 152箇所あった。埋め込みにもキーワード一致にも効くので、必ず潰す。
NUM_SPLIT_RE = re.compile(r"(第)\s*([0-9][0-9\s]*[0-9]|[0-9])\s*(条|章|節|項|号)")
# 本文に紛れ込むページ番号(「- 19 -」「—19—」)。意味を持たないうえ、
# 条文の途中に挟まって文を分断する。
PAGENUM_RE = re.compile(r"^\s*[-—–ー]\s*\d{1,4}\s*[-—–ー]\s*$")


def clean(text: str) -> str:
    """制御文字・ページ番号・数字の桁割れを取り除いて整える。"""
    out = []
    for ch in text:
        o = ord(ch)
        if o < 32 and ch not in "\n\t":
            continue
        out.append(ch)
    s = "".join(out)
    s = unicodedata.normalize("NFKC", s)

    lines = [ln for ln in s.split("\n") if not PAGENUM_RE.match(ln)]
    s = "\n".join(lines)

    # 「第 2 0 条」→「第20条」。数字のあいだの空白だけを消す。
    s = NUM_SPLIT_RE.sub(lambda m: m.group(1) + re.sub(r"\s+", "", m.group(2)) + m.group(3), s)
    # 行末に紛れたページ番号(「…です。 - 19 -」)も除去
    s = re.sub(r"[ \t]*[-—–]\s*\d{1,4}\s*[-—–][ \t]*", " ", s)

    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r"\n{3,}", "\n\n", s)
    return s.strip()


@dataclass
class Page:
    number: int
    text: str


def extract_pages(path: str) -> tuple[list[Page], list[int], list[int]]:
    """(使えるページ, 化けて落としたページ, 目次として落としたページ)

    --layout を付けると pdf_layout(座標ベース)で段組を復元する。

    現時点では既定オフ。理由は実測にある(同一31問での比較):
        素の抽出        R@1 84% / R@3 84% / R@5 90% / ページ81%
        座標で列復元    R@1 81% / R@3 84% / R@5 90% / ページ81%
    文字の欠落は0(実物6資料204ページで確認)だが、列境界が項目名の
    途中に落ちる行があり(「3. 保険期間と｜補償期間 …」)、その行は
    素の抽出より読みにくくなる。作ったから使う、はしない。
    列境界の決め方を直して数字が上回ってから既定にする。
    """
    reader = pypdf.PdfReader(path)
    good: list[Page] = []
    garbled: list[int] = []
    toc: list[int] = []
    for i, page in enumerate(reader.pages, start=1):
        if USE_LAYOUT:
            try:
                raw = pdf_layout.render_page(page)
            except Exception:
                raw = page.extract_text() or ""  # 失敗しても資料を落とさない
        else:
            raw = page.extract_text() or ""
        if len(raw.strip()) < 50:
            continue
        if garble_ratio(raw) > 0.20:
            garbled.append(i)
            continue
        if looks_like_toc(raw):
            toc.append(i)
            continue
        body = clean(raw)
        if len(body) >= MIN_CHARS:
            good.append(Page(i, body))
    return good, garbled, toc


# ---------------------------------------------------------------- 分割


@dataclass
class Block:
    """意味の単位1つ。heading は「第12条(保険金を支払わない場合)」等。"""
    heading: str
    body: str
    page: int


def split_into_blocks(pages: list[Page]) -> list[Block]:
    """条文/見出しを境界にして、ページをまたいで意味の単位へまとめ直す。

    ページ区切りで切らないのが肝。条文はページをまたぐのが普通で、
    ページ単位で切ると条の途中で必ず割れる。
    """
    blocks: list[Block] = []
    cur_heading = ""
    cur_article = ""          # 直近の「第N条(…)」。小見出しはこの下にぶら下げる
    cur_lines: list[str] = []
    cur_page = pages[0].number if pages else 1

    def flush():
        nonlocal cur_lines
        body = "\n".join(cur_lines).strip()
        if len(body) >= MIN_CHARS:
            blocks.append(Block(cur_heading, body, cur_page))
        cur_lines = []

    for page in pages:
        for line in page.text.split("\n"):
            m = ARTICLE_RE.match(line.strip())
            if m:
                flush()
                label = m.group(1)
                title = (m.group(2) or "").strip()
                cur_article = f"{label}({title})" if title else label
                cur_heading = cur_article
                cur_page = page.number
                cur_lines = [line.strip()]
                continue
            if is_heading_line(line):
                flush()
                # 条の中の小見出しなら「第24条 > (2)…」と積む。条を失わない。
                sub = line.strip()
                cur_heading = f"{cur_article} {sub}" if cur_article else sub
                cur_page = page.number
                cur_lines = [sub]
                continue
            cur_lines.append(line)
        # ページ末では flush しない(条文はページをまたぐ)
    flush()
    return blocks


def pack_blocks(blocks: list[Block], source: str) -> list[dict]:
    """ブロックをチャンクへ。長いものだけ分割し、見出しは各断片の先頭に復元する。"""
    chunks: list[dict] = []
    buf: list[Block] = []
    buf_len = 0

    def emit(items: list[Block]):
        if not items:
            return
        head = items[0].heading
        body = "\n".join(b.body for b in items).strip()
        if len(body) < MIN_CHARS:
            return
        text = body
        # 見出しが本文先頭に無ければ復元して付ける(何条の話かを必ず持たせる)
        if head and not text.startswith(head[:8]):
            text = f"【{head}】\n{text}"
        chunks.append({
            "source": source,
            "page": items[0].page,
            "summary": head[:120],
            "keywords": "",
            "full_text": text[:MAX_CHARS],
        })

    for b in blocks:
        if len(b.body) > MAX_CHARS:
            emit(buf); buf, buf_len = [], 0
            # 長い条は項(1. 2. …)で割る。それでも足りなければ文字数で割る。
            parts = re.split(r"\n(?=\s*(?:[(（]?\d+[)）.．]|[①-⑳]))", b.body)
            acc, acc_len = [], 0
            for part in parts:
                if acc_len + len(part) > TARGET_CHARS and acc:
                    emit([Block(b.heading, "\n".join(acc), b.page)])
                    acc, acc_len = [], 0
                acc.append(part); acc_len += len(part)
            emit([Block(b.heading, "\n".join(acc), b.page)])
            continue

        if buf_len + len(b.body) > TARGET_CHARS and buf:
            emit(buf); buf, buf_len = [], 0
        buf.append(b); buf_len += len(b.body)
    emit(buf)
    return chunks



# ---------------------------------------------------------------- 質問の生成

# 見出しから「その資料が確実に答えられる質問」を作る。手で書くと、資料に
# 無いことを聞いてしまい初手で外す。見出し由来なら必ず該当チャンクが当たる。
# 左=見出しに含まれる語 / 右=質問の型。上にあるものほど優先して採用する。
QUESTION_RULES: list[tuple[str, str]] = [
    ("保険金を支払わない",   "{src}で保険金が支払われないのは、どんな場合ですか?"),
    ("引受",                 "{src}の引受にあたって確認すべき点を教えてください。"),
    ("できない",             "{src}で引受できないのは、どんなケースですか?"),
    ("用語の定義",           "{src}の用語の定義を、実務で使う形で説明してください。"),
    ("定義",                 "{h}について、{src}ではどう定義されていますか?"),
    ("支払額",               "{src}の保険金の支払額は、どう計算しますか?"),
    ("支払限度",             "{src}の支払限度額はどうなっていますか?"),
    ("保険料",               "{h}について、{src}の取扱いを教えてください。"),
    ("保険期間",             "{src}の保険期間の取扱いを教えてください。"),
    ("通知義務",             "{src}で通知が必要になるのは、どんなときですか?"),
    ("告知義務",             "{src}の告知義務について教えてください。"),
    ("解除",                 "{src}の契約が解除されるのは、どんなときですか?"),
    ("免責",                 "{src}の免責事項をまとめてください。"),
    ("対象",                 "{h}について、{src}ではどう定めていますか?"),
    ("手続",                 "{src}の手続きの流れを教えてください。"),
    ("とは",                 "{h}について教えてください。"),
]


def strip_label(head: str) -> str:
    """「第12条(保険金を支払わない場合)」→「保険金を支払わない場合」。

    先頭の番号(「14.」「1-2.」)は必ず落とす。落とさないと
    「14. 最低保険料について…」という質問文になる。
    """
    m = re.search(r"[(（]([^)）]+)[)）]", head)
    t = m.group(1).strip() if m else head
    t = re.sub(r"^\s*第\d+条(?:の\d+)?\s*", "", t)
    t = re.sub(r"^\s*(?:[(（]?\d+[)）]|(?:\d+[-－.．]){0,3}\d+[.．)）])\s*", "", t)
    t = re.sub(r"^\s*[◆■●▲]\s*", "", t)
    return t.strip()


def is_good_topic(t: str) -> bool:
    """質問文に埋め込んでよい話題か。

    見出しとして拾った行が、実は本文の一部だったときに
    「に規定する次のいずれかの感染症をいいます。について、…の取扱いを
    教えてください。」という壊れた質問が出る。文の断片を弾く。
    """
    if not t or not (2 <= len(t) <= 30):
        return False
    if any(c in t for c in "。、，"):
        return False
    if t[0] in "にをがはでとやのへも":          # 助詞始まり=文の途中
        return False
    if t.endswith(("います", "ます", "する", "した", "です", "から", "とき")):
        return False
    return True


def build_questions(chunks: list[dict], per_source: int = 6) -> list[dict]:
    """資料ごとに、答えの在り処が確定している質問を作る。"""
    out: list[dict] = []
    by_src: dict[str, list[dict]] = {}
    for c in chunks:
        by_src.setdefault(c["source"], []).append(c)

    for src, items in by_src.items():
        seen_q: set[str] = set()
        picked: list[dict] = []
        short = re.sub(r"\s*[(（].*?[)）]\s*", "", src).strip()
        for rule_kw, tmpl in QUESTION_RULES:
            for c in items:
                head = (c.get("summary") or "").strip()
                if not head or rule_kw not in head:
                    continue
                topic = strip_label(head)
                # {h} を使う型のときだけ話題の質を検査する({src}だけの型は不要)
                if "{h}" in tmpl and not is_good_topic(topic):
                    continue
                q = tmpl.format(src=short, h=topic)
                if q in seen_q:
                    continue
                seen_q.add(q)
                picked.append({"q": q, "source": src, "page": c["page"], "head": head})
                break
            if len(picked) >= per_source:
                break
        out.extend(picked)
    return out


# ---------------------------------------------------------------- 出力


def fnv1a64_hex(s: str) -> str:
    """modUtil.Fnv1a64Hex(src/core/modUtil.bas)と同じ値を返す
    (重複排除キーの互換のため。R12-8で修正: 旧実装はUTF-8バイト列を1バイトずつ
    XOR/乗算していたが、VBA側は文字列をUTF-16コードユニット(AscW値)として
    読み、各コードユニットを「下位バイト→上位バイト」の順に2バイトとして
    別々にXOR/乗算する。ASCIIだけの文字列でもバイト列の粒度が違う(1バイト/字
    vs 2バイト/字)ため、旧実装は全ての非空文字列でVBA側と不一致だった
    (シード資料を本棚へ再取込した際、chunk_idのハッシュ部が一度も一致せず、
    全チャンクが無言で二重登録される実害があった)。

    s.encode("utf-16-le") は文字列を「コードユニットごとに下位バイト→上位
    バイト」の順で並べたバイト列を返すため、VBAの
        code = AscW(ch): if code<0 then code += 65536
        byteLo = code And &HFF: byteHi = (code \\ 256) And &HFF
    と1バイト単位で完全に同じ列になる(サロゲートペア文字も、Pythonの
    utf-16-le エンコードとVBAのAscW走査が同じ2コードユニットを生成するため
    一致する)。"""
    h = 0xCBF29CE484222325
    prime = 0x100000001B3
    mask = 0xFFFFFFFFFFFFFFFF
    data = s.encode("utf-16-le")
    for i in range(0, len(data), 2):
        h ^= data[i]
        h = (h * prime) & mask
        h ^= data[i + 1]
        h = (h * prime) & mask
    return f"{h:016x}"


def normalize_for_hash(s: str) -> str:
    """modUtil.NormalizeForHash と同じ規則で正規化する(R12-8で修正: 旧実装は
    `re.sub(r"\\s+", " ", s)` で改行(LF)も1個の半角スペースへ潰していたが、
    VBA側は「半角スペース/タブの連続だけ」を1個へ圧縮し、改行(LF。CRLF/CRは
    先にLFへ統一)は段落境界の情報として保持したうえでTrim(前後の半角スペース
    のみ除去。改行は前後に残っていても取らない)する規則。この差もFnv1a64Hexと
    同様、実運用のチャンク本文(改行を含む)では常にハッシュ不一致を招いていた。
    """
    t = s.replace("\r\n", "\n").replace("\r", "\n")
    out_chars: list[str] = []
    prev_space = False
    for ch in t:
        if ch == " " or ch == "\t":
            if not prev_space:
                out_chars.append(" ")
            prev_space = True
        else:
            out_chars.append(ch)
            prev_space = False
    return "".join(out_chars).strip(" ")


# ---------------------------------------------------------------- self-test


# modUtil.Fnv1a64Hex / NormalizeForHash をLibreOffice上で実行して実測した
# ゴールデン値(2026-08-01)。src/test/modTestsPure7.bas の
# TestFnvGoldenValues に同じ入力・同じ16進値を複製してVBA側でも固定して
# いるため、この関数のロジックを壊すとどちらのテストも落ちる(=退行が
# 両側から検知できる)。
_SELF_TEST_VECTORS: list[tuple[str, str, str]] = [
    # (ラベル, 入力, 期待hex)
    ("EMPTY", "", "cbf29ce484222325"),
    ("ASCII_HELLO", "hello", "32964f71b2764b97"),
    ("JP_KANJO", "第12条", "5b02a1a0f87d2243"),
    ("LF_MIX", "abc\ndef", "5282a5bd26d7a4f0"),
    ("EMOJI", "OK\U0001F600!", "68ebcbd198394d37"),
    ("SPACES_RAW", "a  b", "d5c496e1f5147176"),
    ("NORM_PIPELINE_RAW", " a   b\t\tc \r\n d ", "6bf429157296e5c6"),
]

# (ラベル, 入力, 期待する正規化後文字列, 正規化後の期待hex)
_SELF_TEST_NORM_VECTORS: list[tuple[str, str, str, str]] = [
    ("SPACES_RAW", "a  b", "a b", "8d862a1a321d76f6"),
    ("NORM_PIPELINE", " a   b\t\tc \r\n d ", "a b c \n d", "2a2ebcad2e5f4f03"),
    ("JP_KANJO", "第12条", "第12条", "5b02a1a0f87d2243"),
]


def self_test() -> bool:
    """fnv1a64_hex/normalize_for_hash がVBA(modUtil.Fnv1a64Hex/
    NormalizeForHash)と同一のゴールデン値を返すことを確認する。
    python3 tools/make_seed_pack.py --self-test で実行する。"""
    ok = True
    for label, raw, expected in _SELF_TEST_VECTORS:
        got = fnv1a64_hex(raw)
        status = "OK" if got == expected else "NG"
        if got != expected:
            ok = False
        print(f"  [{status}] Fnv1a64Hex[{label}] 期待={expected} 実際={got}")

    for label, raw, expected_norm, expected_hash in _SELF_TEST_NORM_VECTORS:
        norm = normalize_for_hash(raw)
        norm_ok = norm == expected_norm
        got_hash = fnv1a64_hex(norm)
        hash_ok = got_hash == expected_hash
        if not (norm_ok and hash_ok):
            ok = False
        status = "OK" if (norm_ok and hash_ok) else "NG"
        print(f"  [{status}] NormalizeForHash[{label}] "
              f"norm期待={expected_norm!r} norm実際={norm!r} "
              f"hash期待={expected_hash} hash実際={got_hash}")

    print()
    if ok:
        print("全件OK: Python側のFnv1a64Hex/NormalizeForHashはVBA(modUtil)と一致します"
              "(src/test/modTestsPure7.bas TestFnvGoldenValues と同じゴールデン値)。")
    else:
        print("!! 不一致があります。tools/make_seed_pack.py の "
              "fnv1a64_hex/normalize_for_hash がVBA(modUtil.Fnv1a64Hex/"
              "NormalizeForHash)と非互換になっています。シード資料の再取込で"
              "chunk_idの重複排除が効かなくなるため修正が必須です。")
    return ok


def write_pack(out_path: str, chunks: list[dict], sources: list[str],
               questions: list[dict] | None = None) -> None:
    wb = Workbook()

    ws = wb.active
    ws.title = "pack_meta"
    ws.append(["key", "value"])
    for k, v in [
        ("pack_format_version", PACK_FORMAT_VERSION),
        ("pack_name", "初期ナレッジ"),
        ("author", "新種保険部 費用・信用グループ"),
        ("created_at", "seed"),
        ("description", " / ".join(sources)),
        ("embed_dim", EMBED_DIM),
        ("app_version", "seed"),
        ("chunk_count", len(chunks)),
    ]:
        ws.append([k, v])
    # 初回ガイドで出す質問。pack_meta の key/value に載せるので、
    # パック形式(3シート)を1バイトも変えずに運べる。
    for i, q in enumerate(questions or [], start=1):
        ws.append([f"question{i}", q["q"]])
        ws.append([f"question{i}_src", f'{q["source"]}|{q["page"]}'])

    wc = wb.create_sheet("pack_chunks")
    wc.append(["chunk_id", "source", "page", "summary", "keywords", "full_text"])
    for i, c in enumerate(chunks, start=1):
        cid = f"bs::{fnv1a64_hex(normalize_for_hash(c['full_text']))}::p{c['page']}::c{i}"
        wc.append([cid, c["source"], c["page"], c["summary"], c["keywords"], c["full_text"]])

    wv = wb.create_sheet("pack_vectors")
    wv.append(["chunk_id", "vector_csv"])   # ベクトルは社内リボンで後入れ

    wb.save(out_path)


def main(argv: list[str]) -> int:
    if "--self-test" in argv:
        print("Fnv1a64Hex/NormalizeForHash セルフテスト "
              "(modUtil.Fnv1a64Hex/NormalizeForHashとの互換性をゴールデン値で照合)")
        return 0 if self_test() else 1

    if len([a for a in argv[1:] if a != "--layout"]) < 2:
        print(__doc__)
        return 2

    global USE_LAYOUT
    args = [a for a in argv[1:] if a != "--layout"]
    USE_LAYOUT = ("--layout" in argv)

    out_path = args[0]
    all_chunks: list[dict] = []
    sources: list[str] = []

    for spec in args[1:]:
        path, _, label = spec.partition(":")
        pages, garbled, toc = extract_pages(path)
        if not pages:
            print(f"  !! 本文を取り出せませんでした: {path}")
            continue
        name = label or path.rsplit("/", 1)[-1].rsplit(".", 1)[0]
        blocks = split_into_blocks(pages)
        chunks = pack_blocks(blocks, name)
        all_chunks.extend(chunks)
        sources.append(name)

        body_chars = sum(len(p.text) for p in pages)
        print(f"  {name}")
        print(f"    使用 {len(pages)}頁 / {body_chars:,}字 → 意味単位 {len(blocks)} → チャンク {len(chunks)}")
        if garbled:
            print(f"    文字化けページを除外: {garbled}")
        if toc:
            print(f"    目次ページを除外: {toc}")

    if not all_chunks:
        print("チャンクが1件も作れませんでした。")
        return 1

    questions = build_questions(all_chunks)
    write_pack(out_path, all_chunks, sources, questions)
    lens = [len(c["full_text"]) for c in all_chunks]
    print()
    print(f"出力: {out_path}")
    print(f"  資料 {len(sources)}件 / チャンク {len(all_chunks)}件 / "
          f"1チャンク平均 {sum(lens)//len(lens)}字 (最大 {max(lens)})")
    print(f"  すぐ押せる質問 {len(questions)}件を同梱(全て資料内に答えがあるもの)")
    print("  ベクトルは未生成。リボンのあるPCで取込→再書き出しして埋めてください。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
