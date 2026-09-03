#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
bin_roundtrip.py - 配布 vbaProject.bin の読み戻し検問(spec_20260903_R35_配布方式転換.md
                   §2-7・§4波2-1。riskconsulting(claude/ai-risk-consulting-spec-review-91o5y5
                   コミット cf628c1)の tools/bin_roundtrip.py を移植)

================================================================================
何を守るゲートか:
    配布方式B(R35)は「モジュールが最初から入った正規の vbaProject.bin を
    ビルドが書き出す」方式である。書き出す側(build/build_mybookshelf.py の
    build_baked_vba_project)が間違っても、書いた側の理屈(build/ovba_write.py)
    で読み返しては何も検査したことにならない。
    本ツールは**配布物(dist/*.xlsm)の側から**bin を開き、

      [1] 各標準モジュールのソースが src/ の .bas/.cls(_vba_src_textの結果。
          Attribute行除去・CRLF・CP932 errors="replace")と**バイト一致**する
      [2] モジュール集合が build/modules.json の台帳(155本の標準モジュール +
          document module {ThisWorkbook, Sheet1})と一致する
      [3] 隠しシート vba_src が存在しない(方式Bはvba_srcを使わない。spec §2-3)
      [4] bin に配布禁止の文字列(build_mybookshelf.FORBIDDEN_BIN_STRINGS=
          "VBProject"/"AddFromString"/"ExecuteExcel4Macro")が現れないこと
          (現れたらFAIL。spec §2-7・§2-8。**解凍した本文で**検査する。
          コメントも対象)。REPORT_BIN_STRINGS("WScript.Shell"/"new:{")は
          件数報告のみ(FAILにしない。R36で機能設計とあわせて撤去)
      [5] 全モジュールの MODULEOFFSET が 0(p-code キャッシュを持たない)こと
      [6] 全 document module のソースが本物のテキストであること(spec §6
          リスク台帳 #6・2026-09-03実Excelで発覚したBLOCKERの再発防止。
          build_mybookshelf.document_module_sanity_errorsを(a)(b)と共有)

    を確かめる。読めない・数えられない・比較できないは**すべて失格**にする
    (「対象が見つからないので検査せず緑」を作らない)。

なぜ oletools.olevba を使うのか:
    [1] の解凍は oletools.olevba(第三者実装)で行う。ビルドが使った
    build/ovba.py の解凍器で読み返すと「自分の圧縮器のバグを自分の解凍器が
    帳消しにする」ため、往復検査としての意味が消える。
    oletools が入っていない場合は**緑にせず** exit 2(環境不備)で止める。

整形規則の値源:
    モジュールソースの整形(Attribute行/.clsヘッダの除去・改行正規化・
    CP932置換)は build/build_mybookshelf.py の `_vba_src_text` /
    `_baked_std_module_bytes` **1実装だけ**を呼ぶ。ここで書き写すと
    ビルドと検問が別々に緩められるため、必ず import して使う。
    禁止文字列の表(FORBIDDEN_BIN_STRINGS / REPORT_BIN_STRINGS)も同じ理由で
    build_mybookshelf.py から import する(二重実装禁止)。

使い方:
    python3 tools/bin_roundtrip.py                 # dist/ の対象ブックを自動検出
    python3 tools/bin_roundtrip.py --book dist/MyBookshelf.xlsm
    exit code: 0 = 全PASS / 1 = 失格 / 2 = 環境不備(oletools不在・ブック不在)
================================================================================
"""

from __future__ import annotations

import argparse
import struct
import sys
import zipfile
from pathlib import Path

TOOLS_DIR = Path(__file__).resolve().parent
REPO_ROOT = TOOLS_DIR.parent
sys.path.insert(0, str(REPO_ROOT / "build"))

import build_mybookshelf as bm   # noqa: E402  (_vba_src_text等・禁止文字列表)
import ovba_write                # noqa: E402  (binの読み戻し)
import ovba                      # noqa: E402  (dirストリームのデコード)

try:
    from oletools.olevba import VBA_Parser
except ImportError:   # 第三者実装が無いと往復検査にならないので緑にしない
    VBA_Parser = None

DEFAULT_BOOKS = (
    REPO_ROOT / "dist" / "MyBookshelf.xlsm",
    REPO_ROOT / "dist" / "MyBookshelf_dev.xlsm",
    REPO_ROOT / "dist" / "MyBookshelf_発行者用.xlsm",
    REPO_ROOT / "dist" / "MyBookshelf_発行者用_dev.xlsm",
)

# 配布方式Bで焼く document module 集合(build_baked_vba_project参照。
# spec §2-11 09-03追記: MyBookshelf は riskconsulting と違い Sheet1 を残す)。
EXPECTED_DOC_MODULES = {"ThisWorkbook", "Sheet1"}


def expected_sources(root: str) -> dict:
    """台帳から {モジュール名: 期待するモジュール本文(bytes、Attribute行除く)} を作る。
    build_baked_vba_project が実際に焼くのと同じ関数(_vba_src_text →
    _baked_std_module_bytes)を通し、Attribute行だけを ovba_write で落とす
    (olevbaが解凍した側もAttribute行を持つため、同じ落とし方で揃える)。"""
    modules = bm.load_manifest(str(REPO_ROOT / "build" / "modules.json"))
    present, _missing = bm.validate_modules(modules, root, allow_missing=True)
    shipped = bm._vba_src_modules(present)
    out = {}
    for m in shipped:
        body = bm._vba_src_text(root, m)
        full_encoded, _replaced = bm._baked_std_module_bytes(m["name"], body)
        out[m["name"]] = ovba_write.strip_attribute_lines(full_encoded)
    return out


def check_book(book: Path) -> list[str]:
    errors: list[str] = []
    print("=" * 78)
    print(f"ブック: {book}")
    print("=" * 78)

    with zipfile.ZipFile(book) as z:
        names = z.namelist()
        if "xl/vbaProject.bin" not in names:
            return [f"{book.name}: xl/vbaProject.bin がありません"]
        vba_bin = z.read("xl/vbaProject.bin")

    # --- [3] vba_src シートの不在 -------------------------------------------
    import openpyxl
    wb = openpyxl.load_workbook(book, read_only=True, keep_links=False)
    sheetnames = list(wb.sheetnames)
    wb.close()
    if "vba_src" in sheetnames:
        errors.append(f"{book.name}: 隠しシート vba_src が残っています"
                      "(配布方式Bでは存在してはいけません)")
    print(f"[3] vba_src シート: {'あり(失格)' if 'vba_src' in sheetnames else 'なし'}"
          f" / シート{len(sheetnames)}枚")

    # --- [1] olevba で解凍して _vba_src_text の結果とバイト比較 ----------------
    got: dict[str, bytes] = {}
    parser = VBA_Parser(str(book))
    try:
        for (_fn, stream, _vba_fn, code) in parser.extract_macros():
            name = stream.split("/")[-1]
            got[name] = code.encode("cp932", errors="replace") \
                if isinstance(code, str) else code
    finally:
        parser.close()
    print(f"[1] olevba が解凍したモジュール: {len(got)}本")

    want = expected_sources(str(REPO_ROOT))
    bin_mods = ovba_write.read_modules(vba_bin)
    doc_names = {n for n, info in bin_mods.items() if info["type"] == "document"}
    if doc_names != EXPECTED_DOC_MODULES:
        errors.append(
            f"{book.name}: document module 集合が {sorted(EXPECTED_DOC_MODULES)} "
            f"と不一致(実際: {sorted(doc_names)})")
    std_got = {n: v for n, v in got.items() if n not in doc_names}

    # --- [6] fail-closed検査(c): 全document moduleのソースが本当にテキストか
    # (spec §6 リスク台帳 #6・2026-09-03実Excelで発覚したBLOCKERの再発防止。
    # build_mybookshelf.document_module_sanity_errorsを(a)(b)と共有する。
    # ここでの読み戻し(ovba_write.read_modules)は配布binの側=baked出力の
    # 側であり、baked出力は全モジュールMODULEOFFSET=0であることを[5]の
    # 検査が別途保証しているので、read_modulesのMODULEOFFSET≠0バグは
    # ここには影響しない)。
    doc_sanity_errors = []
    for nm in sorted(doc_names):
        info = bin_mods.get(nm)
        if info is None:
            continue
        doc_sanity_errors += bm.document_module_sanity_errors(
            nm, info["source"], bm._document_module_expected_base_guid(nm))
    print(f"[6] document module 健全性: "
          f"{'PASS' if not doc_sanity_errors else 'FAIL'}")
    errors += [f"{book.name}: {e}" for e in doc_sanity_errors]

    for name, want_src in sorted(want.items()):
        if name not in std_got:
            errors.append(f"{book.name}: '{name}' が配布binにありません")
            continue
        actual = ovba_write.strip_attribute_lines(std_got[name])
        if actual != want_src:
            errors.append(
                f"{book.name}: '{name}' の本文が src/ と不一致"
                f"(期待{len(want_src)}バイト / 実際{len(actual)}バイト)")
    extra = sorted(set(std_got) - set(want))
    if extra:
        errors.append(f"{book.name}: 台帳に無いモジュールが載っています: {extra}")

    # --- [2] モジュール数と集合 ----------------------------------------------
    print(f"[2] モジュール集合: 台帳{len(want)}本 / bin(document除く){len(std_got)}本"
          f" / document module {sorted(doc_names)}")
    if len(std_got) != len(want):
        errors.append(
            f"{book.name}: モジュール数が台帳と不一致(台帳{len(want)} / bin{len(std_got)})")

    # --- [4] 禁止文字列(FAIL側・件数報告側) ------------------------------------
    hits = bm.forbidden_strings_in_bin(vba_bin)
    print(f"[4] 配布禁止文字列(FAIL側): {hits if hits else 'なし'}")
    if hits:
        errors.append(
            f"{book.name}: vbaProject.bin に配布禁止の文字列があります"
            f"(spec §2-7・§2-8): {', '.join(hits)}")
    report_counts = bm.report_strings_in_bin(vba_bin)
    print("    件数報告(FAILにしない): "
          + ", ".join(f"{w}={n}" for w, n in report_counts.items()))

    # --- [5] MODULEOFFSET=0 ---------------------------------------------------
    dir_dec = ovba.ovba_decompress(ovba.CFBReader(vba_bin).read("dir"))
    offsets = [struct.unpack("<I", body)[0]
               for _o, rid, _s, body in ovba_write.iter_dir_records(dir_dec)
               if rid == ovba_write.REC_MODULEOFFSET]
    bad = [o for o in offsets if o != 0]
    print(f"[5] MODULEOFFSET: {len(offsets)}件すべて{'0' if not bad else '0ではない'}"
          "(p-codeキャッシュ無し)")
    if not offsets:
        errors.append(f"{book.name}: dir に MODULEOFFSET が1件もありません"
                      "(検査が成立していません)")
    if bad:
        errors.append(f"{book.name}: MODULEOFFSET≠0 のモジュールが{len(bad)}件"
                      "(p-codeキャッシュが混入しています)")
    return errors


def main() -> int:
    ap = argparse.ArgumentParser(
        description="配布 vbaProject.bin の読み戻し検問(spec_20260903_R35_配布方式転換.md §2-7)")
    ap.add_argument("--book", help="検査するブック(既定: dist/ の対象ブックを自動検出)")
    args = ap.parse_args()

    if VBA_Parser is None:
        print("ERROR: oletools が import できません(pip install oletools)。"
              "第三者実装で解凍できないと往復検査になりません。", file=sys.stderr)
        return 2

    if args.book:
        p = Path(args.book)
        books = [p if p.is_absolute() else REPO_ROOT / p]
    else:
        books = [p for p in DEFAULT_BOOKS if p.exists()]
    if not books:
        print("ERROR: dist/ にビルド済みブックがありません。"
              "先に python3 build/build_mybookshelf.py --dev を実行してください。",
              file=sys.stderr)
        return 2

    print("=== bin_roundtrip.py (配布binを解凍して src/ とバイト比較) ===")
    errors: list[str] = []
    for b in books:
        if not b.exists():
            print(f"ERROR: ブックが見つかりません: {b}", file=sys.stderr)
            return 2
        errors.extend(check_book(b))
        print()

    print("-" * 78)
    if errors:
        print(f"結果: NG {len(errors)}件")
        for e in errors:
            print(f"  - {e}")
        return 1
    print(f"結果: OK 検査したブック{len(books)}冊 / 6条件"
          "(本文バイト一致・モジュール数と集合・vba_src不在・禁止文字列不在・"
          "MODULEOFFSET=0・document module健全性)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
